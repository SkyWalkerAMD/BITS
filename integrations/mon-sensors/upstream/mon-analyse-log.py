#!/usr/bin/env python3
# -*- coding:utf-8 -*-

import sys
import os
import pandas as pd
import openpyxl
from openpyxl import load_workbook
import xlsxwriter


def read_file(input_filename):
    try:
        with open(input_filename, "r") as fr:
            datas = fr.readlines()[2:]  # 删除前两行数据
    except FileNotFoundError:
        print("无法找到输入文件")
        sys.exit(1)
    return datas


def process_data(datas):
    headers = ["类型", "时间", "当前系统任务", "运行时长", "负载",
               "主频", "cpu温度", "VRM温度", "cpu功耗", "整机功耗", "风扇转速", "e核"]
    data_list = []
    for line in datas:
        data = line.strip().split(",")
        data[5] = data[5].split('/Noa')[0]  # 去除"/Noa"

        # 处理主频列
        if "/" in data[5]:
            main_freq, e_core = data[5].split("/")
            data[5] = main_freq
            data.append(e_core)
        else:
            data.append("")  # 在e核列添加空值

        data_list.append(data)

    df = pd.DataFrame(data_list, columns=headers)
    numeric_columns = ["类型", "运行时长", "负载", "主频",
                       "cpu温度", "VRM温度", "cpu功耗", "整机功耗", "风扇转速", "e核"]
    df[numeric_columns] = df[numeric_columns].replace(
        "Noa", 0).apply(pd.to_numeric, errors="coerce")
    return df


def generate_pivot_table(df):
    sort_order = {'IDIE': 1, 'MLC': 2, 'MBW': 3, 'BC-Result': 4, 'Stress': 5, 'P95-M2': 6,
                  'P95-M1': 7, 'P95-AVX-M2': 8, 'P95-AVX-M1': 9, 'P95-FMA3-M2': 10,
                  'P95-FMA3-M1': 11, 'P95-AVX512-M2': 12, 'P95-AVX512-M1': 13}
    pivot_table = df.pivot_table(index='当前系统任务', values=['负载', '主频', 'cpu温度', 'VRM温度', 'cpu功耗', '整机功耗', '风扇转速'],
                                 aggfunc='mean')
    pivot_table = pivot_table.reindex(
        sorted(pivot_table.index, key=lambda x: sort_order.get(x, 0)))
    pivot_table = pivot_table[['负载', '主频',
                               'cpu温度', 'VRM温度', 'cpu功耗', '整机功耗', '风扇转速']]
    pivot_table = pivot_table.round(2)
    return pivot_table


def generate_chart_data(df, sheet_filename):
    chart_data = {
        '运行时长': ('D', 'red'),
        '负载': ('E', 'orange'),
        '主频': ('F', 'green'),
        'cpu温度': ('G', 'blue'),
        'VRM温度': ('H', 'brown'),
        'cpu功耗': ('I', 'black'),
        '整机功耗': ('J', 'gray'),
        '风扇转速': ('K', 'purple'),
        'e核': ('L', 'magenta')
    }

    chart_data_list = []

    for column, (col_index, color) in chart_data.items():
        data = {
            'column': column,
            'values': f"'{sheet_filename}'!${col_index}$2:${col_index}${len(df) + 1}",
            'color': color
        }
        chart_data_list.append(data)

    return chart_data_list


def generate_charts(workbook, worksheet, chart_data_list):
    chart_index = [22, 38, 54, 71, 87, 103, 119, 135, 152]

    for i, data in enumerate(chart_data_list):
        chart = workbook.add_chart({'type': 'line'})
        chart.add_series({
            'values': data['values'],
            'name': data['column'],
            'line': {'color': data['color']},
        })
        chart_pos = f'M{chart_index[i]}'
        worksheet.insert_chart(chart_pos, chart, {
            'x_offset': 0, 'y_offset': 0, 'x_scale': 1.74, 'y_scale': 1})


def main(input_filename, output_filename, sheet_filename):
    datas = read_file(input_filename)
    df = process_data(datas)

    try:
        # 打开已存在的Excel文件
        excel_file = openpyxl.load_workbook(output_filename)

        # 创建新的工作表
        excel_file.create_sheet(title=sheet_filename, index=0)
        sheet = excel_file[sheet_filename]
        with pd.ExcelWriter(output_filename, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False,
                        sheet_name=sheet_filename, startrow=0)
            pivot_table = generate_pivot_table(df)
            pivot_table.to_excel(
                writer, sheet_name=sheet_filename, startrow=1, startcol=14)

            workbook = writer.book
            worksheet = writer.sheets[sheet_filename]

            chart_data_list = generate_chart_data(df, sheet_filename)
            generate_charts(workbook, worksheet, chart_data_list)

        print("转换完成，并将结果保存到", output_filename)

    except FileNotFoundError:
        with pd.ExcelWriter(output_filename, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_filename)
            pivot_table = generate_pivot_table(df)
            pivot_table.to_excel(
                writer, sheet_name=sheet_filename, startrow=1, startcol=14)

            workbook = writer.book
            worksheet = writer.sheets[sheet_filename]

            chart_data_list = generate_chart_data(df, sheet_filename)
            generate_charts(workbook, worksheet, chart_data_list)

        print("转换完成，并将结果保存到", output_filename)


if __name__ == '__main__':
    input_filename = sys.argv[1]
    output_filename = sys.argv[2]
    sheet_filename = sys.argv[3]
    main(input_filename, output_filename, sheet_filename)
