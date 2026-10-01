"""Fault tests for atomic task claims and real supervised synthetic processes."""
import importlib.util
import json
import os
from pathlib import Path
import shutil
import signal
import shlex
import socket
import subprocess
import sys
import time
from unittest import mock


def make_checks(base):
    class Reliability(base):
        # Inherited cases are loaded once through base, not repeated here.
        @classmethod
        def setUpClass(cls):
            super(Reliability, cls).setUpClass()
            with socket.socket() as sock:
                sock.bind(('127.0.0.1', 0))
                cls.redis_port = sock.getsockname()[1]
            cls.redis = subprocess.Popen(['/usr/bin/redis-server', '--bind', '127.0.0.1',
                '--port', str(cls.redis_port), '--save', '', '--appendonly', 'no'], stdout=subprocess.DEVNULL)
            for i in range(100):
                try:
                    socket.create_connection(('127.0.0.1', cls.redis_port), .1).close()
                    break
                except OSError:
                    time.sleep(.05)

        @classmethod
        def tearDownClass(cls):
            cls.redis.terminate()
            cls.redis.wait(timeout=5)
            super(Reliability, cls).tearDownClass()

        def setUp(self):
            super(Reliability, self).setUp()
            helper = self.app / 'mon-sensors-finish.d'
            for name in ('finish', 'node', 'workload', 'tools_adoption', 'install_guard'):
                sys.modules.pop(name, None)
            sys.path.insert(0, str(helper))
            import node
            self.node = node
            self.queue = node._queue.Queue('127.0.0.1', 'cloud', self.redis_port)
            self.queue.call(0, 'FLUSHALL')
            self.queue.call(0, 'SET', 'cloud', '1')
            for name in ('stress', 'stress-ng'):
                target = self.app / 'bin' / name / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text('#!' + sys.executable + '\nimport time\ntime.sleep(60)\n')
                target.chmod(0o700)

        def tasks(self, names, runtime='2'):
            self.queue.call(1, 'SET', 'IDS', 'CLOUD')
            self.queue.call(1, 'SET', 'DATES', 'batch')
            for name in names:
                self.queue.call(1, 'RPUSH', 'TASKS', name)
                self.queue.call(1, 'SET', name, runtime)

        def node_command(self, check=False):
            args = ['preflight' if check else 'run-queue', '--rdb-server', '127.0.0.1',
                    '--rdb-port', str(self.redis_port), '--log-server', '127.0.0.1:' + str(self.port),
                    '--log-dir', str(self.logs), '--node', 'cloud', '--serial', self.remote_node]
            # Node target is cloud_serial, so allow rsync to create this path.
            return args

        def test_queue_typo_and_duration_preserved_before_claim(self):
            self.tasks(['strss'])
            result = self.cli(*self.node_command(True), good=False)
            self.assertNotEqual(0, result.returncode)
            self.assertEqual('strss', self.queue.call(1, 'LINDEX', 'TASKS', '0'))
            self.assertEqual([], self.module.cases(self.app))
            self.queue.call(1, 'DEL', 'TASKS')
            self.tasks(['stress'], 'bad')
            self.assertNotEqual(0, self.cli(*self.node_command(True), good=False).returncode)
            self.assertEqual('1', self.queue.call(1, 'LLEN', 'TASKS'))

        def test_repeated_names_keep_duration_and_finish_once(self):
            self.tasks(['stress', 'stress'])
            result = self.cli(*self.node_command(), timeout=120)
            state = self.module.cases(self.app)[0]
            self.assertEqual('complete', state['stage'])
            self.assertEqual(2, len(state['steps']))
            self.assertTrue(all(s['execution'] == 'duration_reached' for s in state['steps']))
            self.assertEqual('2', self.queue.call(1, 'GET', 'stress'))
            self.assertEqual('delivered', self.queue.call(1, 'GET', 'STATUS'))

        def test_long_stress_ng_overrides_native_one_day_default(self):
            self.tasks(['stress-ng'], '172800')
            checked = json.loads(self.cli(*self.node_command(True)).stdout)
            argv = checked['tasks'][0]['argv']
            self.assertEqual('172805', argv[argv.index('--timeout') + 1])
            self.assertEqual('1', self.queue.call(1, 'LLEN', 'TASKS'))

        def test_early_exit_preserves_remaining_queue_and_blocks_success(self):
            self.tasks(['stress', 'stress-ng'])
            (self.app / 'bin/stress/stress').write_text('#!/bin/sh\nexit 17\n')
            result = self.cli(*self.node_command(), good=False, timeout=60)
            self.assertNotEqual(0, result.returncode)
            state = self.module.cases(self.app)[0]
            self.assertEqual('early_exit', state['steps'][0]['execution'])
            self.assertEqual(17, state['steps'][0]['exit_code'])
            self.assertTrue(state['steps'][0]['cleanup_confirmed'])
            self.assertEqual('stress-ng', self.queue.call(1, 'LINDEX', 'TASKS', '0'))
            self.assertFalse(Path(state['mon']).with_suffix('.finish.json').exists())
            self.assertIn(b'early_exit', self.cli('status', '--human').stdout)

        def test_atomic_claim_detects_queue_replacement(self):
            self.tasks(['stress'])
            snap = self.queue.snapshot()
            self.queue.call(1, 'SET', 'IDS', 'REPLACED')
            with self.assertRaises(ValueError):
                self.queue.claim(snap, snap['tasks'], 'token')
            self.assertEqual('1', self.queue.call(1, 'LLEN', 'TASKS'))

        def test_empty_queue_returns_without_a_case(self):
            self.assertEqual('idle', json.loads(self.cli(*self.node_command()).stdout)['status'])
            self.assertEqual([], self.module.cases(self.app))

        def test_unrelated_process_survives_workload_stop(self):
            self.tasks(['stress'])
            other = subprocess.Popen(['/usr/bin/sleep', '60'])
            try:
                self.cli(*self.node_command(), timeout=120)
                self.assertIsNone(other.poll())
            finally:
                other.terminate()
                other.wait(timeout=5)

        def test_explicit_tool_copy_preserves_original_ownership(self):
            tool = self.app / 'bin/stress/stress'
            os.chown(str(tool), 201, 200)
            os.chown(str(tool.parent), 201, 200)
            self.tasks(['stress'])
            self.assertNotEqual(0, self.cli(*self.node_command(True), good=False).returncode)
            self.cli('adopt-workloads', '--task', 'stress', '--check')
            self.assertFalse((self.app / '.mon-sensors-workloads').exists())
            self.cli('adopt-workloads', '--task', 'stress')
            self.assertEqual(201, tool.stat().st_uid)
            self.cli(*self.node_command(True))

        def linked_system_tool(self, name='stress'):
            # Match the actual OCRUN layout; the target is synthetic, never a
            # hardware stress program. This runs only in the disposable image.
            tool = self.app / 'bin' / name / name
            target = Path('/usr/bin') / name
            self.assertFalse(target.exists() or target.is_symlink())
            shutil.copyfile(str(tool), str(target))
            target.chmod(0o755)
            self.addCleanup(target.unlink)
            tool.unlink()
            tool.symlink_to(target)
            os.lchown(str(tool), 201, 200)
            os.chown(str(tool.parent), 201, 200)
            os.chown(str(tool.parent.parent), 201, 200)
            return tool, target

        def test_original_system_links_adopt_and_run_without_source_changes(self):
            originals = [self.linked_system_tool(name) for name in ('stress', 'stress-ng')]
            before = [(os.readlink(str(link)), target.read_bytes(), target.stat().st_mode)
                      for link, target in originals]
            args = ('adopt-workloads', '--task', 'stress', '--task', 'stress-ng')
            checked = json.loads(self.cli(*(args + ('--check',))).stdout)
            self.assertEqual('/usr/bin/stress', checked['plans'][0]['files']['stress']['system_target'])
            self.assertFalse((self.app / '.mon-sensors-workloads').exists())
            self.cli(*args)
            for (link, target), (link_text, content, mode) in zip(originals, before):
                copied = self.app / '.mon-sensors-workloads' / target.name / target.name
                self.assertEqual(content, copied.read_bytes())
                self.assertFalse(copied.is_symlink())
                self.assertEqual(0, copied.stat().st_uid)
                self.assertEqual(1, copied.stat().st_nlink)
                self.assertEqual(link_text, os.readlink(str(link)))
                self.assertEqual(201, link.lstat().st_uid)
                self.assertEqual(content, target.read_bytes())
                self.assertEqual(mode, target.stat().st_mode)
            self.tasks(['stress', 'stress-ng'])
            self.cli(*self.node_command(), timeout=120)
            state = self.module.cases(self.app)[0]
            self.assertEqual('complete', state['stage'])
            self.assertEqual(2, len(state['steps']))

        def test_system_link_writable_setid_and_foreign_owner_refused(self):
            link, target = self.linked_system_tool()
            for uid, mode in ((0, 0o777), (0, 0o4755), (201, 0o755)):
                os.chown(str(target), uid, 0)
                target.chmod(mode)
                self.assertNotEqual(0, self.cli('adopt-workloads', '--task', 'stress', good=False).returncode)
                self.assertFalse((self.app / '.mon-sensors-workloads').exists())
            self.assertTrue(link.is_symlink())

        def test_system_link_nonallowlisted_target_refused(self):
            link, target = self.linked_system_tool()
            link.unlink()
            link.symlink_to('/usr/bin/sleep')
            os.lchown(str(link), 201, 200)
            result = self.cli('adopt-workloads', '--task', 'stress', good=False)
            self.assertNotEqual(0, result.returncode)
            self.assertIn(b'Unsupported tool link target', result.stderr)

        def test_system_link_target_under_writable_directory_refused(self):
            link, target = self.linked_system_tool()
            unsafe = self.work / 'writable-system-target'
            unsafe.mkdir()
            copy = unsafe / 'program'
            shutil.copyfile(str(target), str(copy))
            copy.chmod(0o755)
            unsafe.chmod(0o777)
            target.unlink()
            target.symlink_to(copy)
            self.assertNotEqual(0, self.cli('adopt-workloads', '--task', 'stress', good=False).returncode)

        def test_system_link_fifo_and_extra_data_link_refused(self):
            link, target = self.linked_system_tool()
            target.unlink()
            os.mkfifo(str(target), 0o755)
            self.assertNotEqual(0, self.cli('adopt-workloads', '--task', 'stress', good=False, timeout=5).returncode)
            target.unlink()
            target.write_text('#!/bin/sh\nexit 0\n')
            target.chmod(0o755)
            (link.parent / 'extra-data').symlink_to('/etc/hosts')
            self.assertNotEqual(0, self.cli('adopt-workloads', '--task', 'stress', good=False).returncode)
            self.assertFalse((self.app / '.mon-sensors-workloads').exists())

        def test_system_link_retarget_during_read_refused(self):
            link, target = self.linked_system_tool()
            import tools_adoption
            from contextlib import contextmanager
            original = tools_adoption.trusted_executable
            @contextmanager
            def retargeted(path):
                with original(path) as pinned:
                    yield pinned
                link.unlink()
                link.symlink_to('/usr/bin/sleep')
            with mock.patch.object(tools_adoption, 'trusted_executable', retargeted):
                with self.assertRaisesRegex(ValueError, 'Tool link changed'):
                    tools_adoption.adopt(self.app, ['stress'])
            self.assertFalse((self.app / '.mon-sensors-workloads').exists())

        def test_legacy_complete_status_does_not_revoke_prior_verification(self):
            self.tasks(['stress'])
            self.cli(*self.node_command(), timeout=120)
            state = self.module.cases(self.app)[0]
            state['version'] = '0.1.0'
            state.pop('data_quality')
            self.module.persist(self.app, state)
            result = self.cli('status', '--human')
            self.assertIn('旧记录未提供质量字段'.encode('utf-8'), result.stdout)
            self.assertNotIn('数据=尚未核验'.encode('utf-8'), result.stdout)
            self.assertEqual('complete', self.module.cases(self.app)[0]['stage'])

        def test_disk_estimate_refuses_before_claim(self):
            self.tasks(['stress'])
            with mock.patch.object(self.node.shutil, 'disk_usage', return_value=type('Disk', (), {'free': 1})()):
                with self.assertRaises(ValueError):
                    self.node.inspect(self.app, self.queue, self.logs, 'cloud', 'serial')
            self.assertEqual('1', self.queue.call(1, 'LLEN', 'TASKS'))

        def test_scheduler_signal_cleans_workload_and_collector(self):
            self.tasks(['stress', 'stress-ng'], '30')
            child = subprocess.Popen([str(self.entry)] + self.node_command(),
                                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                for i in range(200):
                    files = list((self.app / '.mon-sensors-finish').glob('run-*/step-0.result.json'))
                    if files and json.loads(files[0].read_text()).get('process'):
                        break
                    if child.poll() is not None:
                        self.fail('Scheduler exited before the synthetic task started')
                    time.sleep(.05)
                else:
                    self.fail('No workload process record')
                workload_pid = json.loads(files[0].read_text())['process']['pid']
                child.terminate()
                child.wait(timeout=20)
                result = json.loads(files[0].read_text())
                self.assertTrue(result['cleanup_confirmed'])
                current = self.node.identity(workload_pid)
                self.assertTrue(current is None or current['status'] == 'Z')
                state = self.module.cases(self.app)[0]
                self.assertNotEqual('complete', state['stage'])
                self.assertEqual('stress-ng', self.queue.call(1, 'LINDEX', 'TASKS', '0'))
                self.cli('recover', '--case', state['case'], '--interrupted')
                self.assertEqual('interrupted', self.module.cases(self.app)[0]['outcome'])
            finally:
                if child.poll() is None:
                    child.terminate()
                    child.wait(timeout=20)

        def test_killed_scheduler_guardian_cleans_children(self):
            self.tasks(['stress'], '30')
            child = subprocess.Popen([str(self.entry)] + self.node_command(),
                                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                for i in range(200):
                    files = list((self.app / '.mon-sensors-finish').glob('run-*/step-0.result.json'))
                    if files and json.loads(files[0].read_text()).get('process'):
                        break
                    time.sleep(.05)
                else:
                    self.fail('No workload record')
                child.kill()
                child.wait(timeout=5)
                for i in range(200):
                    result = json.loads(files[0].read_text())
                    if result.get('cleanup_confirmed'):
                        break
                    time.sleep(.05)
                self.assertTrue(result['cleanup_confirmed'])
                self.assertEqual('interrupted', result['execution'])
                state = self.module.cases(self.app)[0]
                self.assertNotEqual('complete', state['stage'])
                self.cli('recover', '--case', state['case'], '--interrupted')
            finally:
                if child.poll() is None:
                    child.terminate()
                    child.wait(timeout=20)

        def test_cloud_large_log_streaming_sheets(self):
            # Exercise a real boundary, preserving every input row and missing cell.
            source = self.logs / 'long.mon'
            with source.open('w') as stream:
                stream.write('#-cloud\n#=type,time,task,uptime,load,mhz,temp,vrm,watts,psu,fan\n')
                for i in range(250001):
                    stream.write('sckocp-v1,20260927_120000,Stress,{},1,4900,55,,120,,\n'.format(i))
            output = source.with_suffix('.xlsx')
            result = subprocess.run(['/usr/local/bin/mon-sensors-report', str(source), str(output)],
                                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
            self.assertEqual(0, result.returncode, result.stderr)
            self.assertEqual(250001, json.loads(result.stdout)['rows'])
            self.assertLess(json.loads(result.stdout)['peak_rss_kib'], 512 * 1024)
            sys.path.insert(0, '/opt/mon-sensors-report/0.2.0/vendor')
            from openpyxl import load_workbook
            workbook = load_workbook(str(output), read_only=True)
            self.assertIn('Monitoring-02', workbook.sheetnames)
            self.assertIsNone(workbook['Monitoring-02']['H2'].value)
            self.assertEqual(250000, workbook['Monitoring-02']['D2'].value)
            workbook.close()

        def test_optional_upload_absent_and_required_upload_failure(self):
            (self.logs / 'result.log').write_text('fixture result\n')
            # Original LOGSVR is a host, not host:port. Pin only the ephemeral
            # loopback port in this test wrapper; exercise real rsync underneath.
            tools = self.work / 'upload-tools'
            tools.mkdir()
            wrapper = tools / 'rsync'
            wrapper.write_text('#!/bin/sh\nexec /usr/bin/rsync --port="$TEST_RSYNC_PORT" "$@"\n')
            wrapper.chmod(0o700)
            environment = dict(os.environ, PATH=str(tools) + ':/usr/bin:/bin', TEST_RSYNC_PORT=str(self.port))
            prefix = 'source {}; LOGPATH={}; HOSTNAME=cloud; MB_SN=fixture; LOGSVR=127.0.0.1; push-log'.format(
                shlex.quote(str(self.app / 'mon-sensors-finish.d/oct-hook.sh')),
                shlex.quote(str(self.logs)))
            good = subprocess.run(['/bin/bash', '-c', prefix], env=environment, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            self.assertEqual(0, good.returncode, good.stderr)
            self.assertIn(b'optional Memtest86', good.stderr)
            environment['TEST_RSYNC_PORT'] = '1'
            bad = subprocess.run(['/bin/bash', '-c', prefix], env=environment,
                                 stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            self.assertNotEqual(0, bad.returncode)

        def test_version_change_stops_before_original_sync(self):
            tools = self.work / 'tools'
            tools.mkdir()
            fake = tools / 'curl'
            fake.write_text('#!/bin/sh\necho MAIN_VERSION=0.9.25\n')
            fake.chmod(0o700)
            (self.app / 'version.txt').write_text('MAIN_VERSION=0.9.24a\n')
            command = 'APPPATH={}; LOGSVR=invalid; source {}; app_check'.format(
                shlex.quote(str(self.app)), shlex.quote(str(self.app / 'mon-sensors-finish.d/hook.sh')))
            result = subprocess.run(['/bin/bash', '-c', command], stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                env=dict(os.environ, PATH=str(tools) + ':/usr/bin:/bin'), timeout=20)
            self.assertNotEqual(0, result.returncode)
            self.assertIn(b'explicit compatible upgrade', result.stderr)

    # Suppress inherited tests to avoid rerunning the entire old fixture suite.
    for name in dir(base):
        if name.startswith('test_') and name not in Reliability.__dict__:
            setattr(Reliability, name, None)
    return Reliability
