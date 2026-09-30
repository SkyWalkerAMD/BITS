"""Offline integration/security checks; run only in disposable cloud containers."""
import fcntl
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import tempfile
import unittest

SRC = Path("/src")
PREFIX = Path("/opt/mon-sensors-report/0.2.0")
COMMAND = "/usr/local/bin/mon-sensors-report"


def run(command, good=True, **kwargs):
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            timeout=60, **kwargs)
    if good and result.returncode:
        raise AssertionError(result.stdout.decode("utf-8", "replace") + result.stderr.decode("utf-8", "replace"))
    return result


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sample(path, task="IDIE", freq=4900):
    path.write_text("#-主机名: fixture provider=sckocp-v1\n"
                    "#=类型,时间,当前系统任务,运行时长,负载,主频,CPU温度(C),VRM温度(C),CPU功耗(W),整机功耗(W),风扇转速\n" +
                    "".join("sckocp-v1,20260926_2100{:02d},{},{},0,{},34,,246.5,,\n".
                            format(i, task, 1000 + i * 2, freq) for i in range(12)), encoding="utf-8")


def app_fixture(app):
    (app / "py").mkdir(parents=True)
    source = (SRC / "integrations/mon-sensors/upstream-0.9.24a/mon-analyse-log.py").read_text(encoding="utf-8")
    source = source.replace("aggfunc='mean')", "aggfunc='mean', dropna=False)").replace(
        "sheet_filename = sys.argv[3]\n", 'sheet_filename = sys.argv[3] if len(sys.argv) > 3 else "Monitoring"\n')
    (app / "py/mon-analyse-log.py").write_text(source, encoding="utf-8")
    (app / "mon-analyse-log").symlink_to("py/mon-analyse-log.py")
    os.lchown(str(app / "mon-analyse-log"), 201, 200)
    for name in ("oct", "ocb", "mon-sensors"):
        shutil.copyfile(str(SRC / "integrations/mon-sensors/upstream-0.9.24a" / name), str(app / name))
        (app / name).chmod(0o755)


class ReportChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(tempfile.mkdtemp(prefix="report-validation-", dir="/root"))
        cls.package = cls.root / "package"
        cls.package.mkdir()
        with tarfile.open(str(SRC / "report-dist/mon-sensors-report-py36-0.2.0.tar.gz")) as bundle:
            bundle.extractall(str(cls.package))
        cls.installer = SRC / "report-dist/mon-sensors-report-py36-0.2.0.run"
        cls.app = Path("/root/ocrun")
        app_fixture(cls.app)
        cls.originals = {name: sha(cls.app / name) for name in
                         ("oct", "ocb", "mon-sensors", "py/mon-analyse-log.py")}
        cls.python3_before = shutil.which("python3")
        baseline = run([sys.executable, "-I", "-S", "-B", "-c",
                        "import importlib.util; assert importlib.util.find_spec('pandas') is None"])
        assert baseline.returncode == 0
        # Reproduce the RPM-installed interpreter layout seen on K6C-165.
        # Only this disposable cloud container gets an extra executable link.
        cls.interpreter = Path(os.path.realpath(sys.executable))
        cls.interpreter_hash = sha(cls.interpreter)
        cls.interpreter_alias = cls.interpreter.parent / ".report-python-hardlink-ci"
        os.link(str(cls.interpreter), str(cls.interpreter_alias))
        assert cls.interpreter.stat().st_nlink >= 2
        sys.path.insert(0, str(cls.package / "payload"))
        import common
        cls.filesystem = common
        try:
            # This is the exact data-file check used by the 0.1.0 launcher.
            common.read(cls.interpreter)
        except ValueError as error:
            assert "hard link" in str(error) and str(cls.interpreter) in str(error)
        else:
            raise AssertionError("The old interpreter check did not reproduce its failure")
        result = run(["bash", str(cls.installer), "--app", str(cls.app), "--check"])
        assert json.loads(result.stdout.decode("utf-8"))["check"] is True
        assert not PREFIX.exists() and not Path(COMMAND).exists()
        assert (cls.app / "mon-analyse-log").is_symlink()
        assert not (cls.app / ".mon-sensors-report-original.json").exists()
        run(["bash", str(cls.installer), "--app", str(cls.app)])
        sys.path.insert(0, str(PREFIX / "vendor"))

    def setUp(self):
        self.work = Path(tempfile.mkdtemp(prefix="case-", dir=str(self.root)))
        self.input = self.work / "input.mon"
        self.output = self.work / "output.xlsx"
        sample(self.input)

    def tearDown(self):
        shutil.rmtree(str(self.work))

    def report(self, good=True, **kwargs):
        return run([COMMAND, str(self.input), str(self.output)], good=good, **kwargs)

    def test_check_and_system_environment_unchanged(self):
        checked = json.loads(run([COMMAND, "--check"]).stdout.decode("utf-8"))
        self.assertEqual("3.6.8", checked["python"])
        self.assertEqual("1.1.5", checked["dependencies"]["pandas"])
        self.assertEqual(self.python3_before, shutil.which("python3"))
        run([sys.executable, "-I", "-S", "-B", "-c", "import importlib.util; assert importlib.util.find_spec('pandas') is None"])

    def test_real_hardlinked_interpreter_installs_and_runs_unchanged(self):
        self.assertGreaterEqual(self.interpreter.stat().st_nlink, 2)
        self.assertEqual(self.interpreter_hash, sha(self.interpreter))
        self.assertEqual(str(self.interpreter), self.filesystem.interpreter_path(sys.executable))
        result = run(["bash", str(self.installer), "--app", str(self.app), "--check"])
        self.assertTrue(json.loads(result.stdout.decode("utf-8"))["check"])
        self.assertEqual("ok", json.loads(self.report().stdout.decode("utf-8"))["status"])

    def test_executable_unsafe_permissions_owner_and_parent_refused(self):
        executable = self.work / "python-fixture"
        executable.write_bytes(b"#!/bin/sh\nexit 0\n")
        for mode in (0o775, 0o757, 0o4755, 0o2755, 0o644):
            executable.chmod(mode)
            with self.assertRaises(ValueError) as raised:
                self.filesystem.interpreter_path(executable)
            self.assertIn(str(executable), str(raised.exception))
        executable.chmod(0o755)
        os.chown(str(executable), 201, 200)
        with self.assertRaises(ValueError) as raised:
            self.filesystem.interpreter_path(executable)
        self.assertIn("uid=201", str(raised.exception))
        os.chown(str(executable), 0, 0)
        self.work.chmod(0o777)
        try:
            with self.assertRaises(ValueError) as raised:
                self.filesystem.interpreter_path(executable)
            self.assertIn(str(self.work), str(raised.exception))
        finally:
            self.work.chmod(0o700)

    def test_data_hardlinks_remain_refused_with_actionable_diagnostic(self):
        alias = self.work / "input-alias.mon"
        os.link(str(self.input), str(alias))
        result = self.report(good=False)
        self.assertNotEqual(0, result.returncode)
        self.assertIn(str(self.input), result.stderr.decode("utf-8"))
        self.assertIn(b"links=2", result.stderr)
        self.assertFalse(self.output.exists())

    def test_managed_install_is_idempotent_and_preserves_collectors(self):
        run(["bash", str(self.installer), "--app", str(self.app)])
        for name, checksum in self.originals.items():
            self.assertEqual(checksum, sha(self.app / name))

    def test_workbook_data_summary_charts_and_blank_columns(self):
        from openpyxl import load_workbook
        source_hash = sha(self.input)
        result = json.loads(self.report().stdout.decode("utf-8"))
        self.assertEqual(12, result["rows"])
        self.assertEqual(source_hash, sha(self.input))
        self.assertEqual(0o600, self.output.stat().st_mode & 0o777)
        workbook = load_workbook(str(self.output))
        sheet = workbook["Monitoring"]
        self.assertEqual("sckocp-v1", sheet["A2"].value)
        self.assertEqual(4900, sheet["F2"].value)
        self.assertEqual(34, sheet["G2"].value)
        self.assertIsNone(sheet["H2"].value)
        self.assertIsNone(sheet["J2"].value)
        self.assertIsNone(sheet["K2"].value)
        self.assertEqual("VRM温度", sheet["S2"].value)
        self.assertIsNone(sheet["S3"].value)
        self.assertEqual(4900, sheet["Q3"].value)
        self.assertEqual(9, len(sheet._charts))
        workbook.close()

    def test_existing_report_backed_up_before_atomic_replacement(self):
        self.report()
        old = self.output.read_bytes()
        sample(self.input, freq=5000)
        result = json.loads(self.report().stdout.decode("utf-8"))
        self.assertEqual(old, Path(result["backup"]).read_bytes())
        self.assertNotEqual(old, self.output.read_bytes())

    def test_formula_and_url_labels_stay_plain_text(self):
        from openpyxl import load_workbook
        for task in ("=1+1", "https://example.invalid/"):
            sample(self.input, task=task)
            self.report()
            workbook = load_workbook(str(self.output))
            cell = workbook["Monitoring"]["C2"]
            self.assertEqual(task, cell.value)
            self.assertNotEqual("f", cell.data_type)
            self.assertIsNone(cell.hyperlink)
            workbook.close()

    def test_output_symlink_and_hardlink_refused(self):
        victim = self.work / "victim"
        victim.write_bytes(b"keep this data")
        self.output.symlink_to(victim)
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.output.unlink()
        os.link(str(victim), str(self.output))
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.assertEqual(b"keep this data", victim.read_bytes())

    def test_input_links_fifo_and_bad_rows_refused(self):
        target = self.work / "real.mon"
        self.input.rename(target)
        self.input.symlink_to(target)
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.input.unlink()
        os.mkfifo(str(self.input), 0o600)
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.input.unlink()
        self.input.write_text("#-fixture\n#=header\nbroken,row\n")
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.assertFalse(self.output.exists())

    def test_untrusted_parent_and_concurrent_output_refused(self):
        self.work.chmod(0o777)
        self.assertNotEqual(0, self.report(good=False).returncode)
        self.work.chmod(0o700)
        lock = self.work / ".output.xlsx.report.lock"
        fd = os.open(str(lock), os.O_CREAT | os.O_RDWR, 0o600)
        try:
            fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
            self.assertNotEqual(0, self.report(good=False).returncode)
        finally:
            os.close(fd)

    def test_pythonpath_injection_ignored(self):
        malicious = self.work / "pandas.py"
        marker = self.work / "injected"
        malicious.write_text("open({!r}, 'w').write('bad')\n".format(str(marker)))
        environment = dict(os.environ, PYTHONPATH=str(self.work), PYTHONHOME="/does-not-exist")
        self.report(env=environment, cwd=str(self.work))
        self.assertFalse(marker.exists())

    def test_modified_analyser_and_entry_refused(self):
        app = self.work / "ocrun"
        app_fixture(app)
        original = (app / "py/mon-analyse-log.py").read_bytes()
        (app / "py/mon-analyse-log.py").write_bytes(original + b"# changed\n")
        self.assertNotEqual(0, run(["bash", str(self.installer), "--app", str(app)], good=False).returncode)
        (app / "py/mon-analyse-log.py").write_bytes(original)
        (app / "mon-analyse-log").unlink()
        (app / "mon-analyse-log").write_text("#!/bin/sh\necho user command\n")
        self.assertNotEqual(0, run(["bash", str(self.installer), "--app", str(app)], good=False).returncode)

    def test_detach_restores_original_symlink_and_ownership(self):
        app = self.work / "ocrun"
        app_fixture(app)
        run(["bash", str(self.installer), "--app", str(app)])
        run(["bash", str(self.installer), "--app", str(app), "--detach", "--check"])
        self.assertFalse((app / "mon-analyse-log").is_symlink())
        run(["bash", str(self.installer), "--app", str(app), "--detach"])
        entry = app / "mon-analyse-log"
        self.assertEqual("py/mon-analyse-log.py", os.readlink(str(entry)))
        self.assertEqual((201, 200), (entry.lstat().st_uid, entry.lstat().st_gid))
        run(["bash", str(self.installer), "--app", str(app)])

    def test_unknown_installed_content_not_silently_repaired(self):
        path = PREFIX / "analyser.py"
        original = path.read_bytes()
        try:
            path.write_bytes(original + b"# modified\n")
            self.assertNotEqual(0, run(["bash", str(self.installer), "--check"], good=False).returncode)
            self.assertNotEqual(original, path.read_bytes())
        finally:
            path.write_bytes(original)

    def test_corrupt_self_extracting_package_refused(self):
        changed = self.work / "corrupt.run"
        changed.write_bytes(self.installer.read_bytes() + b"corrupt")
        result = run(["bash", str(changed), "--check"], good=False)
        self.assertNotEqual(0, result.returncode)
        self.assertIn(b"payload length mismatch", result.stderr)

    def test_renderer_failure_keeps_existing_workbook(self):
        self.report()
        old = self.output.read_bytes()
        result = run([COMMAND, str(self.input), str(self.output), "bad/sheet"], good=False)
        self.assertNotEqual(0, result.returncode)
        self.assertEqual(old, self.output.read_bytes())

    def test_original_oct_analyse_works_without_system_python3(self):
        logs = Path("/root/log")
        logs.mkdir(exist_ok=True)
        sample(logs / "fixture.mon")
        (self.app / "oc.env").write_text(
            "APPPATH=/root/ocrun\nLOGPATH=/root/log\nHOSTNAME=fixture\n"
            "MB_SN=SYNTHETIC\nAPP_DATE=fixture\nMEM_FREE=1024\n")
        run([str(self.app / "oct"), "analyse"], env=dict(os.environ, HOME="/root"))
        self.assertTrue((logs / "fixture-Analyse.xlsx").is_file())
        self.assertEqual(self.python3_before, shutil.which("python3"))


if __name__ == "__main__":
    if os.environ.get("GITHUB_ACTIONS") != "true" or os.geteuid() != 0:
        raise SystemExit("Cloud container only")
    os.umask(0o077)
    result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(ReportChecks))
    summary = {"status": "passed" if result.wasSuccessful() else "failed", "tests_run": result.testsRun,
               "failures": len(result.failures), "errors": len(result.errors),
               "python": sys.version, "os_release": Path("/etc/os-release").read_text(),
               "network": "container started with --network none", "hardware": "synthetic .mon only",
               "interpreter_hardlinks": Path(os.path.realpath(sys.executable)).stat().st_nlink}
    summary_path = Path("/results/report-addon.json")
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    # This is synthetic CI metadata, published by the non-root host runner.
    summary_path.chmod(0o644)
    print(json.dumps(summary))
    sys.exit(not result.wasSuccessful())
